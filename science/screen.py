"""Small deterministic screen tools with explicit approximations and provenance."""
from __future__ import annotations

import datetime as dt
import hashlib
import os
import re
import time
from functools import lru_cache
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DATA = Path(os.environ.get("NEXT_EXPERIMENT_DATA", ROOT / "data/countmatrix.xlsx" if (ROOT / "data/countmatrix.xlsx").exists() else ROOT.parent / "research/decision-followup/countmatrix.xlsx"))
COLUMNS = ["premscreening_D0_rep1", "premscreening_D0_rep2", "IL6m_rep1", "IL6m_rep2", "IL6p_rep1", "IL6p_rep2"]
CONDITIONS = {"without_il6": [0, 1], "with_il6": [2, 3]}
PARAMETERS = {
    "analysis_version": "0.1.0", "normalization": "total library depth to geometric-mean depth, then non-targeting median centering per contrast",
    "pseudocount_normalized_reads": 1.0, "baseline_min_each_repeat": 20,
    "min_eligible_guides": 3, "effect_threshold_log2": 0.5,
    "max_repeat_difference_log2": 1.0, "min_positive_guide_fraction": 0.6,
    "gene_summary": "median of guide mean log2 enrichments across two repeats",
    "baseline_filter": "same guides pass baseline-count threshold in both D0 repeats",
}
LIMITATIONS = [
    "Exploratory full-data reanalysis; no independent held-out evaluation or measured agent advantage.",
    "Approximate log2 count enrichment, not MAGeCK MLE beta scores, a reproduced publication analysis, p-values or FDR.",
    "Positive enrichment means relative persistence/expansion of guide-bearing cells; it does not prove rejuvenation, preserved cell identity, safety or longevity.",
    "Negative enrichment may reflect impaired cell fitness; count depletion does not uniquely establish a senescence mechanism.",
    "Two measured repeats only; guide leave-one-out is a sensitivity check, not new biological replication.",
    "Operational thresholds are explicit heuristics, not calibrated probabilities or gene-level statistical significance.",
    "All pooled measurements already exist; no wet-lab experiment or cost savings were demonstrated.",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def guide_gene(guide: str) -> str:
    if re.fullmatch(r"non-targeting_\d+", guide):
        return "non-targeting"
    m = re.fullmatch(r"(.+)_[+-]_\d+", guide)
    if not m:
        raise ValueError(f"Unrecognized sgRNA identifier: {guide!r}")
    return m.group(1)


def resolve_gene(guide: str, cell) -> tuple[str, dict | None]:
    derived = guide_gene(guide)
    if isinstance(cell, (dt.date, dt.datetime)):
        return derived, {"guide": guide, "original_gene_cell": cell.isoformat(), "resolved_gene": derived, "reason": "Excel date-typed label; exact original symbol recovered from sgRNA identifier"}
    if not isinstance(cell, str) or not cell.strip():
        raise ValueError(f"Missing/non-text gene mapping: {guide!r}")
    label = cell.strip()
    if (label == "non-targeting") != (derived == "non-targeting"):
        raise ValueError(f"Non-targeting label and guide identifier disagree: {guide!r}, {label!r}")
    if label != derived:
        # Explicit gene cell is retained; a locus suffix may differ from the gene name.
        return label, {"guide": guide, "original_gene_cell": label, "resolved_gene": label, "guide_prefix": derived, "reason": "Text label differs from guide prefix; retain explicit Gene cell, do not infer an alias"}
    return label, None


def normalize_counts(counts: np.ndarray) -> tuple[np.ndarray, np.ndarray, float]:
    counts = np.asarray(counts, dtype=float)
    if counts.ndim != 2 or counts.shape[1] != 6 or not np.isfinite(counts).all() or (counts < 0).any():
        raise ValueError("Counts must be a finite nonnegative N x 6 matrix")
    totals = counts.sum(axis=0)
    if (totals <= 0).any():
        raise ValueError("All six libraries must have positive total counts")
    reference = float(np.exp(np.mean(np.log(totals))))
    return counts * (reference / totals), totals, reference


def guide_enrichments(normalized: np.ndarray, control_mask: np.ndarray, pseudocount: float = 1.0) -> tuple[np.ndarray, np.ndarray]:
    if pseudocount <= 0 or not np.isfinite(pseudocount):
        raise ValueError("Pseudocount must be finite and positive")
    control_mask = np.asarray(control_mask, dtype=bool)
    if control_mask.shape != (len(normalized),) or not control_mask.any():
        raise ValueError("Verified non-targeting guides are required for control centering")
    posts = normalized[:, [2, 3, 4, 5]]
    baselines = normalized[:, [0, 1, 0, 1]]
    raw = np.log2((posts + pseudocount) / (baselines + pseudocount))
    centers = np.median(raw[control_mask], axis=0)
    return raw - centers, centers


@dataclass
class Screen:
    counts: np.ndarray
    guides: np.ndarray
    genes: np.ndarray
    path: Path
    digest: str
    mapping_notes: list[dict]

    def __post_init__(self):
        self.normalized, self.totals, self.reference = normalize_counts(self.counts)
        self.controls = self.genes == "non-targeting"
        self.effects, self.control_centers = guide_enrichments(self.normalized, self.controls)
        self.eligible = np.all(self.counts[:, :2] >= PARAMETERS["baseline_min_each_repeat"], axis=1)

    @classmethod
    def load(cls, path: Path = DEFAULT_DATA):
        path = Path(path).resolve()
        wb = load_workbook(path, read_only=True, data_only=True)
        if wb.sheetnames != ["CountMatrix"]:
            raise ValueError(f"Unexpected sheets: {wb.sheetnames}")
        rows = wb.active.iter_rows(values_only=True)
        if list(next(rows)) != ["sgRNA", "Gene", *COLUMNS]:
            raise ValueError("Unexpected count matrix column schema")
        guides, genes, counts, notes = [], [], [], []
        for row in rows:
            guide = str(row[0])
            gene, note = resolve_gene(guide, row[1])
            guides.append(guide); genes.append(gene); counts.append(row[2:])
            if note:
                notes.append(note)
        wb.close()
        if len(set(guides)) != len(guides):
            raise ValueError("Duplicate sgRNA identifiers")
        a = np.asarray(counts, dtype=float)
        if not np.isfinite(a).all() or not np.equal(a, np.floor(a)).all():
            raise ValueError("Counts must be observed nonnegative integers without missing entries")
        return cls(a, np.array(guides), np.array(genes), path, sha256(path), notes)

    def qc(self) -> dict:
        counts_per_gene = pd.Series(self.genes).value_counts()
        controls = self.effects[self.controls]
        return {
            "guide_rows": int(len(self.genes)), "distinct_gene_labels_including_controls": int(len(counts_per_gene)),
            "target_gene_labels": int(len(counts_per_gene) - 1),
            "library_totals": dict(zip(COLUMNS, self.totals.astype(int).tolist())),
            "normalization_reference_depth": self.reference,
            "zero_count_guides_by_library": dict(zip(COLUMNS, (self.counts == 0).sum(axis=0).astype(int).tolist())),
            "baseline_eligible_guides": int(self.eligible.sum()),
            "control_inventory": {"validated_label": "non-targeting", "guide_id_pattern": "non-targeting_<digits>", "count": int(self.controls.sum()), "eligible_count": int((self.controls & self.eligible).sum()), "do_not_treat_as_control": ["CTRL", "NONO", "KNTC1"]},
            "control_raw_centers_log2": self.control_centers.tolist(),
            "control_centered_guide_quantiles_log2": {str(q): np.quantile(controls, q, axis=0).tolist() for q in (0.05, 0.5, 0.95)},
            "mapping_notes": self.mapping_notes,
            "date_label_repairs": sum("Excel date" in n["reason"] for n in self.mapping_notes),
            "text_label_prefix_differences_retained": sum("Text label" in n["reason"] for n in self.mapping_notes),
            "guide_count_distribution": {str(k): int(v) for k, v in counts_per_gene.drop("non-targeting").value_counts().sort_index().items()},
        }

    def candidate(self, gene: str) -> dict:
        if gene == "non-targeting":
            raise ValueError("Non-targeting guides are controls, not candidate genes")
        indices = np.flatnonzero(self.genes == gene)
        if not len(indices):
            raise ValueError(f"Unknown exact gene symbol: {gene!r}")
        selected = indices[self.eligible[indices]]
        result = {"gene": gene, "total_guides": int(len(indices)), "eligible_guides": int(len(selected)), "conditions": {}, "missing_endpoints": ["cell identity", "senescence phenotype for this complete candidate set", "safety", "human longevity"]}
        for condition, cols in CONDITIONS.items():
            result["conditions"][condition] = summarize_condition(self.effects[selected][:, cols], self.guides[selected].tolist())
        result["context"] = summarize_context(self.effects[selected], self.guides[selected].tolist())
        result["guides"] = [{"sgRNA": str(self.guides[i]), "eligible": bool(self.eligible[i]), "baseline_counts": self.counts[i, :2].astype(int).tolist(), "centered_log2_enrichment": self.effects[i].tolist()} for i in indices]
        result["baseline_count_sensitivity"] = self.count_sensitivity(indices)
        return result

    def count_sensitivity(self, indices: np.ndarray) -> list[dict]:
        """Predeclared filter variants; record all, never select the best result."""
        outputs = []
        for threshold in (10, 20, 50):
            selected = indices[np.all(self.counts[indices, :2] >= threshold, axis=1)]
            conditions = {}
            for condition, cols in CONDITIONS.items():
                summary = summarize_condition(self.effects[selected][:, cols], self.guides[selected].tolist())
                conditions[condition] = {"score": summary["score"], "repeat_scores": summary["repeat_scores"], "loo_range": [summary["leave_one_out"]["score_min"], summary["leave_one_out"]["score_max"]] if summary["leave_one_out"] else None}
            outputs.append({"baseline_min_each_repeat": threshold, "eligible_guides": len(selected), "conditions": conditions})
        return outputs

    def table(self) -> pd.DataFrame:
        # Vectorized aggregate avoids thousands of repeated scans of the guide matrix.
        effects = pd.DataFrame(self.effects[self.eligible], columns=["m1", "m2", "p1", "p2"])
        effects.insert(0, "gene", self.genes[self.eligible])
        effects["m_mean"] = effects[["m1", "m2"]].mean(axis=1)
        effects["p_mean"] = effects[["p1", "p2"]].mean(axis=1)
        effects["ctx1"] = effects["p1"] - effects["m1"]
        effects["ctx2"] = effects["p2"] - effects["m2"]
        effects["ctx_mean"] = effects[["ctx1", "ctx2"]].mean(axis=1)
        grouped = effects.groupby("gene", sort=True)
        result = grouped[["m1", "m2", "p1", "p2", "m_mean", "p_mean", "ctx1", "ctx2", "ctx_mean"]].median()
        result["eligible_guides"] = grouped.size()
        result["total_guides"] = pd.Series(self.genes).value_counts()
        result = result.drop("non-targeting", errors="ignore")
        for short in ("m", "p"):
            result[f"{short}_repeat_sign_agreement"] = (np.sign(result[f"{short}1"]) == np.sign(result[f"{short}2"]))
            result[f"{short}_positive_guide_fraction"] = grouped[f"{short}_mean"].apply(lambda x: float((x > 0).mean()))
            # LOO on guide means keeps the estimator identical to the full-gene score.
            result[f"{short}_loo_min"] = grouped[f"{short}_mean"].apply(lambda x: loo_bounds(x.to_numpy())[0])
            result[f"{short}_loo_max"] = grouped[f"{short}_mean"].apply(lambda x: loo_bounds(x.to_numpy())[1])
            result[f"{short}_fixed_checklist_pass"] = (
                (result["eligible_guides"] >= PARAMETERS["min_eligible_guides"])
                & (result[f"{short}1"] >= PARAMETERS["effect_threshold_log2"])
                & (result[f"{short}2"] >= PARAMETERS["effect_threshold_log2"])
                & (result[f"{short}_loo_min"] >= PARAMETERS["effect_threshold_log2"])
                & ((result[f"{short}1"] - result[f"{short}2"]).abs() <= PARAMETERS["max_repeat_difference_log2"])
                & (result[f"{short}_positive_guide_fraction"] >= PARAMETERS["min_positive_guide_fraction"])
            )
        result["context_repeat_sign_agreement"] = np.sign(result["ctx1"]) == np.sign(result["ctx2"])
        return result.reset_index()


def loo_bounds(values: np.ndarray) -> tuple[float, float]:
    values = np.asarray(values, dtype=float)
    if len(values) < 2:
        return float("nan"), float("nan")
    medians = [float(np.median(np.delete(values, i))) for i in range(len(values))]
    return min(medians), max(medians)


def summarize_condition(values: np.ndarray, guides: list[str]) -> dict:
    if len(values) == 0:
        return {"status": "insufficient_baseline_coverage", "eligible_guides": 0, "score": None, "repeat_scores": None, "leave_one_out": None, "fixed_checklist_pass": False}
    means = values.mean(axis=1)
    score = float(np.median(means))
    repeats = np.median(values, axis=0)
    loo = [{"omitted_guide": guide, "score": float(np.median(np.delete(means, i)))} for i, guide in enumerate(guides)] if len(values) >= 2 else []
    low = min((r["score"] for r in loo), default=None)
    high = max((r["score"] for r in loo), default=None)
    positive_fraction = float((means > 0).mean())
    passes = bool(len(values) >= PARAMETERS["min_eligible_guides"] and (repeats >= PARAMETERS["effect_threshold_log2"]).all() and low is not None and low >= PARAMETERS["effect_threshold_log2"] and abs(repeats[1] - repeats[0]) <= PARAMETERS["max_repeat_difference_log2"] and positive_fraction >= PARAMETERS["min_positive_guide_fraction"])
    return {"status": "computed", "eligible_guides": len(values), "score": score, "repeat_scores": repeats.tolist(), "repeat_difference": float(repeats[1] - repeats[0]), "repeat_sign_agreement": bool(np.sign(repeats[0]) == np.sign(repeats[1])), "positive_guide_fraction": positive_fraction, "leave_one_out": {"status": "computed" if loo else "insufficient_guides", "score_min": low, "score_max": high, "sign_reversal": bool(any(np.sign(r["score"]) != np.sign(score) for r in loo)), "threshold_reversal": bool(any((r["score"] >= PARAMETERS["effect_threshold_log2"]) != (score >= PARAMETERS["effect_threshold_log2"]) for r in loo)), "omissions": loo}, "fixed_checklist_pass": passes}


def summarize_context(values: np.ndarray, guides: list[str]) -> dict:
    if len(values) == 0:
        return {"status": "insufficient_baseline_coverage", "score": None, "repeat_scores": None}
    differences = values[:, [2, 3]] - values[:, [0, 1]]
    repeats = np.median(differences, axis=0)
    return {"status": "computed", "score": float(np.median(differences.mean(axis=1))), "repeat_scores": repeats.tolist(), "repeat_sign_agreement": bool(np.sign(repeats[0]) == np.sign(repeats[1])), "meaning": "Positive contrast indicates stronger guide enrichment with IL6 than without IL6", "dependence": "Within-guide contrasts share D0 and subtract before aggregation; D0 cancels algebraically. Context contrasts are not independent samples.", "leave_one_out": summarize_condition(differences, guides)["leave_one_out"]}


@lru_cache(maxsize=2)
def _cached_screen(path: str, mtime_ns: int) -> Screen:
    return Screen.load(Path(path))


def run_tool(operation: str, gene: str | None = None, data_path: str | Path = DEFAULT_DATA) -> dict:
    """Import-friendly, compact Omnigent boundary; no arbitrary code or model calls.

    These development tools expose both repeats. Do not use this API as a final
    evaluation selector. Full correction ledger is generated by the baseline CLI.
    """
    started = time.perf_counter()
    allowed = {"qc", "candidate", "guide-check", "context", "baseline", "triage", "dependence"}
    if operation not in allowed:
        return {"status": "failed", "tool": operation, "error": "Operation is not allowlisted"}
    limitations = list(LIMITATIONS)
    try:
        path = Path(data_path).resolve()
        screen = _cached_screen(str(path), path.stat().st_mtime_ns)
        if operation == "dependence":
            if not gene:
                raise ValueError("Exact gene symbol is required")
            from science.dependence import decision_dependence
            observations = decision_dependence(screen, gene)
            limitations.append("Items marked observed=false are hypothetical probes of the fixed rule (one added guide); they are not measurements and must not be reported as observed results.")
        elif operation == "qc":
            observations = screen.qc()
            observations["mapping_note_count"] = len(observations["mapping_notes"])
            observations["mapping_note_examples"] = observations.pop("mapping_notes")[:5]
        elif operation == "baseline":
            table = screen.table()
            observations = {"scope": "exploratory_full_dataset_both_repeats", "scored_gene_count": len(table), "top_candidates": {}}
            for short, condition in (("m", "without_il6"), ("p", "with_il6")):
                cols = ["gene", "eligible_guides", f"{short}_mean", f"{short}1", f"{short}2", f"{short}_loo_min", f"{short}_loo_max", f"{short}_fixed_checklist_pass"]
                observations["top_candidates"][condition] = table.loc[table.eligible_guides >= PARAMETERS["min_eligible_guides"], cols].sort_values(f"{short}_mean", ascending=False).head(5).to_dict("records")
        elif operation == "triage":
            if not gene or gene == "non-targeting":
                raise ValueError("Exact target gene symbol is required")
            from benchmark.triage import evaluate
            indices = np.flatnonzero(screen.genes == gene)
            if not len(indices):
                raise ValueError(f"Unknown exact gene symbol: {gene!r}")
            selected = indices[screen.eligible[indices]]
            observations = {"gene": gene, "eligible_guides": len(selected), "conditions": {
                name: evaluate(screen.effects[selected][:, cols], adaptive=True)
                for name, cols in CONDITIONS.items()
            }, "meaning": "Stops after a decisive checklist failure; exact same heuristic decision as exhaustive checking. A failure requests more evidence, not biological rejection."}
        else:
            if not gene:
                raise ValueError("Exact gene symbol is required")
            candidate = screen.candidate(gene)
            if operation == "candidate":
                observations = candidate
            elif operation == "context":
                observations = {"gene": gene, **candidate["context"]}
            else:
                observations = {"gene": gene, "eligible_guides": candidate["eligible_guides"], "conditions": {k: {"score": v["score"], "leave_one_out": v["leave_one_out"], "fixed_checklist_pass": v["fixed_checklist_pass"]} for k,v in candidate["conditions"].items()}, "baseline_count_sensitivity": candidate["baseline_count_sensitivity"]}
        return {"status": "computed", "tool": operation, "evidence_type": "computed_from_observed_counts", "input_hashes": {"countmatrix.xlsx": screen.digest}, "input_kind": "counterfactual_evaluation" if screen.path.name.startswith("COUNTERFACTUAL-") else "observed_public_data", "parameters": PARAMETERS, "observations": observations, "artifact_paths": [], "limitations": limitations, "duration_seconds": round(time.perf_counter()-started, 4), "resource_use": {"model_calls": 0, "api_cost_usd": 0}}
    except (ValueError, OSError) as e:
        return {"status": "failed", "tool": operation, "error": str(e), "duration_seconds": round(time.perf_counter()-started, 4)}
